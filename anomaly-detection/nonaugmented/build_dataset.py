# -*- coding: utf-8 -*-
"""Constroi o dataset rotulado NAO-AUMENTADO para a avaliacao real de RQ1.

Unidade = uma linha de consumo tal como o host a propos, antes de correcao.
Label   = a linha foi corrigida pelo processo manual (Quantidade Alterada
          ou Referencia Alterada no relatorio de diferencas).

REGRA CRITICA: a quantidade ja corrigida do PHC nunca entra como feature.
Para as linhas sinalizadas usa-se `Qtt host` (valor pre-correcao); para as
nao sinalizadas host == phc por definicao.
"""
import openpyxl, json, datetime, csv, os
from collections import defaultdict, Counter

# Diretoria com os exports reais do ERP (nao versionados -- ver .gitignore).
# Definir DATA_DIR para apontar para a copia local; ver README.
DATA_DIR = os.environ.get('DATA_DIR', './data')

U = os.path.join(DATA_DIR, '01_DATA_CONSUMOS') + os.sep
LIST = U + '_tmp_listagem_copy.xlsx'
BANQ = U + 'Consumos_Banquetes_Detalhe_Consolidado.xlsx'
DIFF = U + 'Consumos_Relatorio_de_Diferencas_de_Consumos_Host_Vs_PHC_1805a2806.xlsx'

def d(s):
    return datetime.datetime.strptime(s, '%d.%m.%Y').date()

# ---------------------------------------------------------------- listagem
wb = openpyxl.load_workbook(LIST, read_only=True, data_only=True)
rows = list(wb.worksheets[0].iter_rows(values_only=True))
lines = []
for r in rows[4:]:
    if r[1] in (None, ''):
        continue
    lines.append(dict(stream=str(r[1]).strip(), doc=str(r[2]).strip(), date=d(str(r[3]).strip()),
                      ref=str(r[4] or '').strip(), desc=str(r[5] or '').strip(),
                      qty=float(r[6] or 0), unit_cost=float(r[7] or 0),
                      value=float(r[8] or 0), wh=str(r[9] or '').strip(), src='listagem'))
# volume diario por fluxo (folha Resumo por Consumo)
ws = wb['Resumo por Consumo']
dayvol = {}
for r in list(ws.iter_rows(values_only=True))[3:]:
    if r[0] in (None, '') or r[1] in (None, '') or r[2] in (None, ''):
        continue
    dayvol[(str(r[1]).strip(), d(str(r[2]).strip()))] = (float(r[4] or 0), float(r[5] or 0))
wb.close()

# ---------------------------------------------------------------- banquetes
wb = openpyxl.load_workbook(BANQ, read_only=True, data_only=True)
for r in list(wb['Banquete_Detalhe'].iter_rows(values_only=True))[1:]:
    if r[0] in (None, ''):
        continue
    dt = datetime.date.fromisoformat(str(r[0])[:10])
    lines.append(dict(stream='Eventos Consumo', doc=str(r[3]).strip(), date=dt,
                      ref=str(r[5]).strip(), desc=str(r[6] or '').strip(),
                      qty=float(r[7] or 0), unit_cost=float(r[9] or 0),
                      value=float(r[10] or 0), wh=str(r[11] or '').strip(),
                      src='banquetes', event=str(r[4]).strip()))
wb.close()
print('linhas totais: %d' % len(lines))
print('  por fluxo:', dict(Counter(l['stream'] for l in lines)))

# volume diario para o fluxo de eventos (nao vem em folha de resumo)
ev = defaultdict(lambda: [0.0, 0.0])
for l in lines:
    if l['stream'] == 'Eventos Consumo':
        v = ev[l['date']]; v[0] += l['qty']; v[1] += l['value']
for dt, v in ev.items():
    dayvol[('Eventos Consumo', dt)] = tuple(v)

# ---------------------------------------------------------------- diferencas
wb = openpyxl.load_workbook(DIFF, read_only=True, data_only=True)
diffs = []
for r in list(wb.worksheets[0].iter_rows(values_only=True))[5:]:
    if r[1] in (None, ''):
        continue
    diffs.append(dict(stream=str(r[1]).strip(), doc=str(r[2]).strip(), date=d(str(r[3]).strip()),
                      refphc=str(r[4] or '').strip(), refhost=str(r[5] or '').strip(),
                      qphc=float(r[8] or 0), qhost=float(r[9] or 0), nota=str(r[10] or '').strip()))
wb.close()
corr = [x for x in diffs if x['nota'] in ('Quantidade Alterada', 'Referência Alterada')]
print('discrepancias corrigidas (nao apagadas): %d' % len(corr))
print('  apagadas (fora de ambito):', sum(1 for x in diffs if x['nota'] == 'Linha Apagada'))

# ---------------------------------------------------------------- emparelhar
idx = defaultdict(list)
for i, l in enumerate(lines):
    idx[(l['stream'], l['doc'], l['ref'])].append(i)

matched, ambiguous, unmatched = 0, 0, []
for x in corr:
    key = (x['stream'], x['doc'], x['refphc'] or x['refhost'])
    cand = [i for i in idx.get(key, []) if lines[i].get('label') is None]
    if not cand:
        unmatched.append(x); continue
    # desempate pela quantidade registada == Qtt phc
    exact = [i for i in cand if abs(lines[i]['qty'] - x['qphc']) < 1e-6]
    if len(exact) >= 1:
        pick = exact[0]
        if len(exact) > 1: ambiguous += 1
    else:
        pick = min(cand, key=lambda i: abs(lines[i]['qty'] - x['qphc']))
        ambiguous += 1
    l = lines[pick]
    l['label'] = 1
    l['nota'] = x['nota']
    l['qty_phc'] = l['qty']
    l['qty'] = x['qhost']                      # quantidade PRE-correcao
    if x['nota'] == 'Referência Alterada' and x['refhost']:
        l['ref'] = x['refhost']                 # referencia pre-correcao
    matched += 1

for l in lines:
    if l.get('label') is None:
        l['label'] = 0
        l['qty_phc'] = l['qty']                 # host == phc

print('emparelhadas: %d de %d  (desempates por quantidade: %d)' % (matched, len(corr), ambiguous))
if unmatched:
    print('  NAO emparelhadas:', len(unmatched), [(u['stream'], u['doc'], u['refphc']) for u in unmatched[:5]])
pos = sum(l['label'] for l in lines)
print('positivos: %d de %d  (taxa %.2f%%)' % (pos, len(lines), 100 * pos / len(lines)))
print('  por fluxo:', {k: (sum(1 for l in lines if l['stream'] == k),
                           sum(l['label'] for l in lines if l['stream'] == k))
                       for k in sorted(set(l['stream'] for l in lines))})

json.dump([{**l, 'date': l['date'].isoformat()} for l in lines], open('lines.json', 'w'))
json.dump({f'{k[0]}|{k[1].isoformat()}': v for k, v in dayvol.items()}, open('dayvol.json', 'w'))
print('\ngravado lines.json (%d) e dayvol.json (%d)' % (len(lines), len(dayvol)))
